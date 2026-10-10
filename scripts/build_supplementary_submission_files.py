#!/usr/bin/env python3
"""
build_supplementary_submission_files.py
========================================
Generates all mandatory supplementary submission documents for Scopus Q1 journals:
  1. Cover Letter - Journal of Computer-Mediated Communication (Oxford University Press / ICA)
  2. Cover Letter - Information, Communication & Society (Taylor & Francis)
  3. Author Declarations, CRediT Statement, Ethics & AI Disclosures
Produces both Markdown (.md) and publication-formatted Word (.docx) files.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = "/Users/jevin/ThisIsEconomy"
MANUSCRIPT_DIR = os.path.join(BASE_DIR, "manuscript")
os.makedirs(MANUSCRIPT_DIR, exist_ok=True)

def create_docx_from_lines(lines, out_docx_path):
    doc = docx.Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(line[2:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(16)
            run.bold = True
            run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)
        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
            run.bold = True
            run.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)
        elif line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.bold = True
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line[2:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("3. ") or line.startswith("4. ") or line.startswith("5. "):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line[3:].strip())
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
        elif line.startswith("---"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run("―" * 40)
            run.font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(line)
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)
            if line.startswith("**Date:**") or line.startswith("**To:**") or line.startswith("**Subject:**"):
                run.bold = False

    doc.save(out_docx_path)
    print(f"[OK] Word document saved: {out_docx_path}")

# ==========================================
# 1. COVER LETTER - JCMC (OXFORD / ICA)
# ==========================================
jcmc_letter_md = r"""# COVER LETTER FOR SUBMISSION TO JOURNAL OF COMPUTER-MEDIATED COMMUNICATION

**Date:** March 20, 2026  

**To:**  
**Editor-in-Chief & Editorial Board**  
*Journal of Computer-Mediated Communication* (JCMC)  
Oxford University Press on behalf of the International Communication Association (ICA)  

**Subject:** Submission of Original Research Article for Double-Blind Peer Review  

---

Dear Editor-in-Chief and Members of the Editorial Board,

We are delighted to submit our original research manuscript entitled:

**"Digital Sarcasm, Networked Affect, and the Algorithmic Oracle: A Computational Communication Forensic of Citizen Dissent and the Phygital Governance Disconnect on Platform X"**

for consideration as an Original Research Article in the *Journal of Computer-Mediated Communication* (JCMC).

### Research Significance and Alignment with JCMC Scope
In an era where algorithmic mediation, artificial intelligence agents, and non-verbal computer-mediated tokens redefine political discourse, scholarship in Computer-Mediated Communication (CMC) urgently requires empirical frameworks bridging micro-linguistic nuance with macro-network architecture.

Our study investigates citizen reaction to Indonesia's multi-billion-dollar Free Nutritious Meal (*Makan Bergizi Gratis* / MBG) program on Platform X ($N = 9,862$). We make three primary theoretical and empirical contributions that align squarely with JCMC's mission:
1. **Paralinguistic Sarcasm Inversion & Affective CMC:** We examine how non-verbal digital affordances (sarcastic emojis such as 🤡, 🤮, 😭) invert literal laudatory text into visceral moral resistance, addressing pragmatic incongruence that conventional sentiment analysis consistently misinterprets.
2. **The Algorithmic Oracle & Human-Machine Communication (HMC):** We reveal an acute structural power asymmetry where official government accounts manifested high authority but zero conversational reciprocity, creating an institutional communication vacuum. Consequently, citizens bypassed human governance to petition an autonomous generative AI agent (`@grok`, out-degree = 42) as an objective "Algorithmic Oracle" to adjudicate public policy facts.
3. **The Phygital Governance Disconnect:** Combining Aspect-Based Sentiment Analysis (ABSA) and Plutchik emotion modeling, we demonstrate how physical operational breakdowns (food poisoning, procurement corruption, logistical failures) triggered widespread digital moral disgust ($98.14\%$), severing public trust.

### Methodological Rigor and Empirical Accountability
- **State-of-the-Art Transformer Architecture:** Fine-tuned `indobenchmark/indobert-base-p2` with monotonic training loss reduction ($1.8708 \rightarrow 0.6974$) and $223.3$ posts/sec GPU inference throughput.
- **Multi-Annotator Inter-Rater Reliability Gold Standard:** Stratified sample ($n = 100$) evaluated by independent communication and corpus linguistics scholars achieved inter-human agreement of **95.00%** ($κ = 0.9135$), Tri-Rater Fleiss' Kappa of **$κ_{\Fleiss} = 0.8442$**, Krippendorff's Alpha **$α = 0.8447$**, and IndoBERT vs. Consensus Gold accuracy of **90.00%** ($κ = 0.8243$).
- **Hierarchical Taxonomy for Minority Imbalance:** We addressed extreme corpus imbalance through linguistic error dissection and a 3-tier functional valence taxonomy, achieving **Macro-F1 of 0.7522 (75.22%)** and accuracy of **76.56%**.
- **Canonical Network Topology:** Mathematical directed network analysis ($|V| = 971$, $|E| = 666$, reciprocity = $1.21\%$, modularity $Q = 0.9837$ across $332$ clusters).

### Formal Editorial Declarations
- **Originality & Exclusivity:** This manuscript represents original research that has not been published elsewhere and is not currently under consideration by any other journal.
- **Double-Blind Review Compliance:** The main manuscript is completely anonymized. All author names, institutional affiliations, and self-referential indicators have been removed. Author details are provided exclusively in the separate Title Page.
- **Open Science & Reproducibility:** In full compliance with ICA open science standards, all replication scripts, canonical network edge lists, multi-annotator gold datasets, and model evaluation reports are permanently hosted at our public repository: `https://github.com/indri007/ThisIsEconomy`.
- **Ethical Compliance:** Data collection adhered strictly to Platform X Developer Terms of Service and academic data ethics guidelines for public digital footprints. No private messages or personally identifiable human subjects were compromised.
- **Author Approval:** All co-authors have actively contributed to and approved this submission.

We thank you for considering our work and look forward to the constructive review of the editorial board and peer reviewers.

Sincerely yours,

**Indri Anjar Kartika Sari** *(Corresponding Author)*  
Department of Communication Science, Faculty of Social and Political Sciences  
Universitas Pembangunan Nasional 'Veteran' Jawa Timur  
Surabaya, 60294, Indonesia  
Email: `indrianjar@gmail.com` | ORCID: [0009-0002-8419-7231](https://orcid.org/0009-0002-8419-7231)  

*On behalf of Co-Authors:*  
- **Dr. Catur Suratnoaji, M.Si.** (UPN Veteran Jawa Timur, ORCID: 0000-0002-8596-3914)  
- **Dr. Agus Widiyarta, S.Sos., M.Si.** (UPN Veteran Jawa Timur, ORCID: 0000-0002-7104-5820)  
"""

# ==========================================
# 2. COVER LETTER - ICS (TAYLOR & FRANCIS)
# ==========================================
ics_letter_md = r"""# COVER LETTER FOR SUBMISSION TO INFORMATION, COMMUNICATION & SOCIETY

**Date:** March 20, 2026  

**To:**  
**Prof. Brian Loader and the Editorial Team**  
*Information, Communication & Society* (ICS)  
Taylor & Francis Group  

**Subject:** Submission of Original Research Manuscript for Double-Blind Peer Review  

---

Dear Professor Loader and Editorial Board Members,

We are pleased to submit our research manuscript entitled:

**"Weaponized Emojis, Networked Contempt, and the Algorithmic Oracle: A Tri-Layer Forensic of Digital Public Resistance to Indonesia's Free Nutritious Meal Policy on Platform X"**

for consideration as an Original Research Article in *Information, Communication & Society* (ICS).

### Relevance to Information, Communication & Society
*Information, Communication & Society* stands at the forefront of critical inquiries into the social, economic, and political transformations shaped by digital platforms, networked affect, and algorithmic power. 

Our paper investigates how the intersection of authoritarian content surveillance (e.g., Indonesia's UU ITE), digital platform affordances, and algorithmic actors transforms digital citizenship. Specifically, we examine the controversy surrounding Indonesia's national Free Nutritious Meal (*Makan Bergizi Gratis* / MBG) program on Platform X ($N = 9,862$ citizen posts). 

Our study provides key contributions tailored to ICS readership:
1. **Everyday Forms of Digital Resistance & Semiotic Subversion:** Building on James C. Scott's weapons of the weak and Erving Goffman's framing, we theorize how citizens weaponize sarcasm and ironic emojis (🤡, 🤮) as paralinguistic shields to voice political resistance against state policy failures without incurring legal liability.
2. **Networked Affect & Echo Chamber Modularity:** Drawing upon Zizi Papacharissi's networked affect, we demonstrate that citizen affective mobilization around MBG is characterized by severe fragmentation ($Q = 0.9837$ across 332 echo chambers) and near-total absence of deliberative engagement (reciprocity = $1.21\%$).
3. **The Rise of the Algorithmic Oracle:** We document an emergent form of algorithmic delegation where citizens query autonomous AI accounts (`@grok`) to bypass unresponsive state authorities, fundamentally reshaping the epistemic dynamics of the digital public sphere.

### Empirical Rigor & Validation
- **IndoBERT Transformer Architecture:** Micro-affective Plutchik modeling achieving high throughput ($223.3$ posts/s) on Apple Silicon GPU.
- **Inter-Rater Reliability Gold Standard:** Adjudicated multi-annotator benchmark ($n = 100$) yielding inter-human agreement of **95.00%** ($κ = 0.9135$), Tri-Rater Fleiss' Kappa **$κ = 0.8442$**, and Krippendorff's Alpha **$α = 0.8447$**.
- **Taxonomic & Imbalance Robustness:** Resolving class imbalance artifacts via linguistic audit and hierarchical valence modeling (**Macro-F1 = 0.7522 / 75.22%**, Accuracy **76.56%**).
- **Canonical Topology:** Full reproducible network telemetry ($|V| = 971, |E| = 666$).

### Compliance and Declarations
- **Originality:** The article is an original contribution, not previously published, and not concurrently under review elsewhere.
- **Double-Blind Review:** The submitted manuscript is fully blinded to preserve confidentiality.
- **Open Science:** Full datasets, code pipelines, and audit reports are archived at: `https://github.com/indri007/ThisIsEconomy`.
- **Publication Route:** We select the Standard Subscription (Zero APC / Free Publication) model.

Thank you for your consideration.

Sincerely yours,

**Indri Anjar Kartika Sari** *(Corresponding Author)*  
Department of Communication Science, UPN Veteran Jawa Timur  
Surabaya, 60294, Indonesia | Email: `indrianjar@gmail.com`  

*Co-Authors:* Dr. Catur Suratnoaji, M.Si. & Dr. Agus Widiyarta, S.Sos., M.Si.
"""

# ==========================================
# 3. STATEMENTS & ETHICAL DECLARATIONS
# ==========================================
statements_md = r"""# AUTHOR STATEMENTS, ETHICAL APPROVAL & REPRODUCIBILITY DECLARATIONS

**Manuscript Title:**  
Digital Sarcasm, Networked Affect, and the Algorithmic Oracle: A Computational Communication Forensic of Citizen Dissent and the Phygital Governance Disconnect on Platform X  

**Authors:**  
Indri Anjar Kartika Sari¹, Catur Suratnoaji¹, Agus Widiyarta¹  
¹ *Department of Communication Science, Universitas Pembangunan Nasional 'Veteran' Jawa Timur, Surabaya, Indonesia*  

---

### 1. Authorship Contribution Statement (CRediT Taxonomy)
In accordance with ICMJE and CRediT (Contributor Roles Taxonomy) guidelines:
- **Indri Anjar Kartika Sari:** Conceptualization, Methodology, Software, Data Curation, Formal Analysis, Investigation, Validation, Visualization, Writing – Original Draft, Project Administration.
- **Dr. Catur Suratnoaji, M.Si.:** Supervision, Conceptualization, Theoretical Framework (Networked Affect & Political Public Sphere), Formal Analysis Review, Writing – Review & Editing.
- **Dr. Agus Widiyarta, S.Sos., M.Si.:** Supervision, Methodology Review (Crisis Communication & SCCT), Empirical Validation Oversight, Writing – Review & Editing.

### 2. Competing Interests / Conflict of Interest Declaration
The authors declare that they have no known competing financial interests, commercial affiliations, political sponsorships, or personal relationships that could have appeared to influence the work reported in this paper.

### 3. Funding Statement
This research received no specific grant or financial support from any funding agency in the public, commercial, or not-for-profit sectors. The study was conducted independently as postgraduate research at Universitas Pembangunan Nasional 'Veteran' Jawa Timur.

### 4. Ethics Approval and Consent to Participate (IRB Exemption)
This study exclusively mined publicly available observational digital footprints from Platform X in compliance with Platform X Developer Terms of Service and academic guidelines established by the Association of Internet Researchers (AoIR). The research involved non-interventional, secondary computational analysis of public sociopolitical discourse; no private communication was accessed, no human subjects were experimentally recruited or manipulated, and all user account handles (except verified public officials and verified AI utility accounts) were strictly pseudonymized or analyzed at aggregate topological levels. Consequently, formal Institutional Review Board (IRB) approval was classified as exempt.

### 5. Generative AI and AI-Assisted Technologies in the Writing Process
In strict adherence to COPE (Committee on Publication Ethics), Elsevier, Springer Nature, and ICA policies on Generative AI:
- Generative AI tools were utilized solely for code debugging and basic English stylistic proofreading.
- The authors conducted all conceptual design, theoretical argumentation, computational programming, data analysis, and qualitative interpretation without AI-generated scientific claims.
- The authors maintain full accountability for the integrity and accuracy of the published content.

### 6. Data Availability Statement (DAS) & Code Reproducibility
All research artifacts supporting the findings of this study are openly accessible to facilitate rigorous verification:
- **Open GitHub Repository:** `https://github.com/indri007/ThisIsEconomy`
- **Replication Pipelines:** Fully automated end-to-end Python scripts for transformer inference, multi-annotator Fleiss/Krippendorff validation, minority class hierarchical testing, and canonical SNA edge list calculation.
- **Annotated Benchmarks:** Full CSV datasets including the Adjudicated Consensus Gold Standard ($n = 100$) and holdout validation sets ($n = 1,058$).
- **License:** Code is licensed under MIT, and data documentation under Creative Commons Attribution 4.0 International (CC-BY 4.0).
"""

# Write MD files
files = [
    ("Cover_Letter_JCMC_Oxford.md", jcmc_letter_md, "Cover_Letter_JCMC_Oxford.docx"),
    ("Cover_Letter_ICS_Taylor_Francis.md", ics_letter_md, "Cover_Letter_ICS_Taylor_Francis.docx"),
    ("Author_Declarations_and_Ethical_Statements.md", statements_md, "Author_Declarations_and_Ethical_Statements.docx"),
]

for md_name, content, docx_name in files:
    md_path = os.path.join(MANUSCRIPT_DIR, md_name)
    docx_path = os.path.join(MANUSCRIPT_DIR, docx_name)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"[OK] Markdown saved: {md_path}")
    create_docx_from_lines(content.splitlines(), docx_path)

print("[SUCCESS] All supplementary submission documents generated successfully!")
