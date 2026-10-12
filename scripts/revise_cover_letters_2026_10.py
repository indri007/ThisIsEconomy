#!/usr/bin/env python3
"""
Revisi cover letter ICS & JCMC (12 Okt 2026): menghapus klaim reliabilitas
antarpenilai / gold standard, menampilkan metrik utama 6 kelas lebih dulu,
lalu membangun ulang .docx dari .md (tanpa sisa tanda ** atau $ di Word).
"""
import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt

MS = Path(__file__).resolve().parent.parent / "manuscript"

VALID_BULLET = ("- **Transparent Validation Scope:** Emotion metrics are reported against silver-standard "
                "labels on a group-aware holdout set ($n = 1,058$): accuracy **75.99%**, macro-F1 **0.4535**, "
                "weighted-F1 **0.7434** across the six emotion classes present in the data. The manuscript states "
                "explicitly that no independent human gold standard has yet been constructed and that Anger and "
                "Sadness are not recognised by the model.")
TAXO_ICS_OLD = ("- **Taxonomic & Imbalance Robustness:** Resolving class imbalance artifacts via linguistic audit "
                "and hierarchical valence modeling (**Macro-F1 = 0.7522 / 75.22%**, Accuracy **76.56%**).")
TAXO_JCMC_OLD = ("- **Hierarchical Taxonomy for Minority Imbalance:** We addressed extreme corpus imbalance through "
                 "linguistic error dissection and a 3-tier functional valence taxonomy, achieving **Macro-F1 of "
                 "0.7522 (75.22%)** and accuracy of **76.56%**.")
TAXO_NEW = ("- **Supplementary Valence Analysis:** As a post-hoc complement to the granular results, a three-tier "
            "valence grouping (negative, positive, neutral) yields macro-F1 **0.7522** and accuracy **76.56%**.")


def must(t, old, new, where):
    if old not in t:
        sys.exit(f"[STOP] tidak ditemukan di {where}: {old[:60]!r}")
    return t.replace(old, new)


def revise_ics(t):
    s = t.index("- **Inter-Rater Reliability Gold Standard:**")
    e = t.index("\n", s)
    t = t[:s] + VALID_BULLET + t[e:]
    return must(t, TAXO_ICS_OLD, TAXO_NEW, "ICS taxonomy")


def revise_jcmc(t):
    s = t.index("- **Multi-Annotator Inter-Rater Reliability Gold Standard:**")
    e = t.index("\n", s)
    t = t[:s] + VALID_BULLET + t[e:]
    t = must(t, TAXO_JCMC_OLD, TAXO_NEW, "JCMC taxonomy")
    return must(t, "multi-annotator gold datasets,", "silver-labelled emotion datasets,", "JCMC open science")


def clean_inline(s):
    s = re.sub(r"\$([^$]+)\$", lambda m: m.group(1), s)
    s = s.replace(r"\rightarrow", "→").replace(r"κ_{\Fleiss}", "κ").replace("\\", "")
    s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", s)
    return s


def add_runs(p, text, size=11):
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", clean_inline(text)):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1])
        else:
            r = p.add_run(tok)
        r.font.name = "Times New Roman"; r.font.size = Pt(size)


def md_to_docx(md, out):
    d = docx.Document()
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
    for line in md.split("\n"):
        raw = line.rstrip()
        if not raw.strip() or raw.strip() == "---":
            continue
        if raw.startswith("# "):
            p = d.add_paragraph(); add_runs(p, f"**{raw[2:].strip()}**", 13)
        elif raw.startswith("### "):
            p = d.add_paragraph(); add_runs(p, f"**{raw[4:].strip()}**", 11.5)
        elif re.match(r"^\s*- ", raw):
            p = d.add_paragraph(style="List Bullet"); add_runs(p, re.sub(r"^\s*- ", "", raw))
        elif re.match(r"^\d+\. ", raw):
            p = d.add_paragraph(style="List Number"); add_runs(p, re.sub(r"^\d+\. ", "", raw))
        else:
            p = d.add_paragraph(); add_runs(p, raw.strip())
        p.paragraph_format.space_after = Pt(6)
    d.save(str(out))


if __name__ == "__main__":
    for name, fn in (("Cover_Letter_ICS_Taylor_Francis", revise_ics), ("Cover_Letter_JCMC_Oxford", revise_jcmc)):
        md_path = MS / f"{name}.md"
        t = md_path.read_text(encoding="utf-8")
        if "Transparent Validation Scope" not in t:
            t = fn(t)
            md_path.write_text(t, encoding="utf-8")
        md_to_docx(t, MS / f"{name}.docx")
        print(f"[OK] {name}.md + .docx")
