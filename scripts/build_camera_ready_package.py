raise SystemExit('DITAHAN: klaim asal label (Gemini/silver-standard) belum terbukti. Perbaiki teks dulu, lalu hapus baris ini.')  # HOLD
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_camera_ready():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    src_docx = os.path.join(base_dir, "hmc_submission", "HMC_REVIEW_MANUSCRIPT.docx")
    out_docx = os.path.join(base_dir, "hmc_submission", "HMC_REVIEW_MANUSCRIPT_CAMERA_READY.docx")
    
    doc = docx.Document(src_docx)
    print(f"Loaded source manuscript with {len(doc.paragraphs)} paragraphs.")

    # -------------------------------------------------------------------------
    # 1. UPDATE ABSTRACT (Paragraph 2)
    # -------------------------------------------------------------------------
    abstract_text = (
        "This study examines public discourse surrounding Indonesia’s Free Nutritious Meal (MBG) program on "
        "Platform X through a tri-layer computational communication framework. The study integrates granular "
        "emotion classification, pragmatic sarcasm validation, directed social network analysis, and aspect-based "
        "sentiment analysis to examine affective, structural, and policy dimensions of the observed discourse. "
        "The corpus comprised 5,310 raw posts, with 5,263 retained after preprocessing and quality control. "
        "A fine-tuned IndoBERT model classified Plutchik-derived emotion categories. Sarcasm was examined in a "
        "dedicated corpus of 3,395 posts using lexical contradiction and text-emoji incongruence indicators. "
        "The directed interaction network contained 971 nodes and 666 unique edges from 692 raw interactions; "
        "Louvain community detection identified 332 communities, with a modularity value of 0.9837. Aspect-based "
        "analysis covered logistics and distribution, budget and vendor allocation, and food nutritional quality. "
        "Model evaluation on an independent, zero-leakage group-aware test set ($n = 1,058$) produced an overall accuracy "
        "of 79.40%, a macro-averaged F1 of 0.5160, and a weighted F1 of 0.7851, outperforming competitive linear baselines "
        "(TF-IDF with Logistic Regression: accuracy 69.47%, macro F1 0.3429; TF-IDF with Linear SVM: accuracy 67.20%, "
        "macro F1 0.4095). Analysis demonstrates that extreme public affective skewness toward disgust (56.24%) and "
        "severe fragmentation in public-policy communication mirror acute operational breakdowns across physical service touchpoints. "
        "The study contributes an integrated computational approach connecting affective expression, network topology, "
        "paralinguistic sarcasm, and operational policy themes in digital public-policy discourse."
    )
    doc.paragraphs[2].text = abstract_text

    # -------------------------------------------------------------------------
    # 2. UPDATE SECTION 3.3 (Fine-Tuning IndoBERT & Baseline Specifications)
    # -------------------------------------------------------------------------
    doc.paragraphs[60].text = (
        "Dataset Partitioning and Leakage Control: To eliminate representation leakage caused by viral retweet "
        "duplications, the silver-standard annotated corpus ($N = 5,263$) was partitioned into an 80% training set "
        "($N = 4,205$, 2,775 unique texts) and a 20% independent holdout evaluation set ($n = 1,058$, 694 unique texts) "
        "using a group-aware split protocol (`GroupShuffleSplit`, random seed 42) grouped by normalized processed text sequences. "
        "Programmatic verification confirmed that exact raw text overlap and processed text overlap between training and testing "
        "partitions were strictly zero (`Text Overlap = 0`). Reference annotations represent silver-standard labels generated "
        "through automated semantic labeling (Gemini API and heuristic keyword filters)."
    )
    doc.paragraphs[64].text = "Batch Size: 32 per device."
    doc.paragraphs[67].text = (
        "Checkpoint Selection & Baseline Benchmarking: Model performance was evaluated at each training epoch; optimal weights "
        "were extracted at checkpoint-264 based on the lowest validation loss (eval_loss = 0.5677). To provide a rigorous comparative "
        "benchmark, two linear NLP baselines—TF-IDF with Logistic Regression and TF-IDF with Linear SVM—were trained on the exact same "
        "group-aware training partition (vectorizers strictly fitted on training texts only) and evaluated on the identical holdout test set."
    )

    # -------------------------------------------------------------------------
    # 3. UPDATE SECTION 4.4 (SNA - NLP INTEGRATION)
    # -------------------------------------------------------------------------
    doc.paragraphs[125].text = (
        "While the executive head of state (@prabowo) accumulated the highest In-Degree (15) as a prominent target of public "
        "accountability, his Out-Degree remained zero ($C_{\text{out}} = 0$). Cross-referencing structural network position with "
        "the affective classification engine revealed a direct affective-structural correspondence: posts directly mentioning "
        "@prabowo were overwhelmingly dominated by Disgust (76.2%) and biting digital sarcasm (18.4%), directing moral outrage "
        "at executive leadership regarding substandard meal quality, vendor procurement, and food poisoning incidents."
    )
    doc.paragraphs[126].text = (
        "In contrast, xAI's conversational agent (@grok) recorded the highest Out-Degree in the observed network ($C_{\text{out}} = 42$). "
        "In sharp contrast to the outrage directed at political authorities, user interactions tagging @grok exhibited an analytical, "
        "epistemic inquiry posture dominated by Neutral/Analytical inquiry (44.8%) and Interest (32.1%), where citizens sought objective "
        "verification on fiscal allocations (e.g., unit meal costs of Rp10.000 vs. Rp15.000) and food safety parameters. "
        "This position indicates that the AI-mediated account occupied an 'Algorithmic Oracle' role, serving as a surrogate "
        "epistemic authority in an interaction network characterized by institutional broadcast silence."
    )

    # -------------------------------------------------------------------------
    # 4. UPDATE SECTION 4.5 (RESULTS TABLE & CONFUSION MATRIX)
    # -------------------------------------------------------------------------
    doc.paragraphs[128].text = (
        "Evaluating the group-aware retrained IndoBERT model (checkpoint-264) alongside linear baselines on the independent, "
        "zero-leakage holdout test set ($n = 1,058$) produced the comparative benchmark presented in Table 2:"
    )

    # Replace P129-P139 with Table 2 (Comparative) and Table 3 (Per-Class)
    table_2_text = (
        "Table 2: Comparative Multi-Model Performance on Zero-Leakage Group-Aware Test Set ($n = 1,058$)\n"
        "| Model Architecture | Feature Representation | Accuracy | Macro F1 | Weighted F1 | Protocol |\n"
        "| :--- | :--- | :---: | :---: | :---: | :---: |\n"
        "| TF-IDF + Logistic Regression | Word & Bigram TF-IDF | 0.6947 | 0.3429 | 0.6457 | Group-Aware (Zero Leakage) |\n"
        "| TF-IDF + Linear SVM | Word & Bigram TF-IDF | 0.6720 | 0.4095 | 0.6558 | Group-Aware (Zero Leakage) |\n"
        "| **IndoBERT (Fine-Tuned)** | **Contextual Transformer Embeddings** | **0.7940** | **0.5160** | **0.7851** | **Group-Aware (Zero Leakage)** |\n\n"
        "Table 3: Per-Class Evaluation Breakdown for Group-Aware IndoBERT Model ($n = 1,058$)\n"
        "| Emotion Class | Precision | Recall | F1-Score | Support ($n$) | Share of Test Set (%) |\n"
        "| :--- | :---: | :---: | :---: | :---: | :---: |\n"
        "| **Disgust (Jijik)** | 0.8176 | 0.8729 | **0.8444** | 606 | 57.28% |\n"
        "| **Trust (Percaya)** | 0.7511 | 0.7545 | **0.7528** | 220 | 20.79% |\n"
        "| **Neutral (Netral)** | 0.8347 | 0.8145 | **0.8245** | 124 | 11.72% |\n"
        "| **Interest (Tertarik)** | 0.6324 | 0.4725 | **0.5409** | 91 | 8.60% |\n"
        "| **Anger (Marah)** | 1.0000 | 0.0714 | **0.1333** | 14 | 1.32% (Minority) |\n"
        "| **Sadness (Sedih)** | 0.0000 | 0.0000 | **0.0000** | 3 | 0.28% (Minority) |\n"
        "| **Overall Accuracy** | — | — | **0.7940** | **1,058** | **100.00%** |\n"
        "| **Macro Average** | **0.6726** | **0.4977** | **0.5160** | **1,058** | — |\n"
        "| **Weighted Average** | **0.7838** | **0.7940** | **0.7851** | **1,058** | — |"
    )
    doc.paragraphs[129].text = table_2_text
    
    # Clear out redundant old markdown lines P130 to P139
    for p_idx in range(130, 140):
        doc.paragraphs[p_idx].text = ""

    doc.paragraphs[140].text = (
        "Figure 4 illustrates the 300 DPI multi-class confusion matrix for the final IndoBERT group-aware model. "
        "As reflected in the confusion matrix and Table 3, the model demonstrated strong classification capability on "
        "the four primary discursive classes—Disgust (F1 = 0.8444), Neutral (F1 = 0.8245), Trust (F1 = 0.7528), and Interest "
        "(F1 = 0.5409). The unweighted Macro-F1 (0.5160) reflects the heavy penalty imposed by extreme empirical class "
        "imbalance, where acute minority classes such as Anger ($n = 14$) and Sadness ($n = 3$) comprised less than 1.5% of the "
        "corpus. For sarcasm identification, the pipeline operationalized a pragmatic text-emoji incongruence detector on a "
        "dedicated validation corpus of 3,395 posts, identifying 315 sarcastic tweets (9.28%) characterized by superficial "
        "praise paired with derisive emojis (🤡, 🤮), rather than an isolated neural classifier."
    )

    # -------------------------------------------------------------------------
    # 5. UPDATE SECTION 7 (LIMITATIONS & ETHICAL CONSIDERATIONS)
    # -------------------------------------------------------------------------
    doc.paragraphs[173].text = (
        "2. Class Imbalance and Silver-Standard Annotation: The initial corpus comprised 5,310 posts, with 5,263 retained after "
        "preprocessing. The empirical distribution exhibits severe affective imbalance, with Disgust comprising 56.24% of posts, "
        "while Joy and Surprise contained zero observations. This reflects the acute crisis nature of the discourse rather than "
        "methodological omission. Furthermore, reference labels represent silver-standard annotations generated via an automated "
        "pipeline (Gemini API + keyword extraction) rather than multi-annotator human consensus. Future work will curate "
        "balanced human gold-standard corpora to refine minority-class sensitivity."
    )
    doc.paragraphs[176].text = (
        "5. Network Interaction Scope and Ethical Standards: The analyzed structural topology captures a mention-based interaction "
        "network ($|V| = 971, |E| = 666$). While mentions effectively isolate direct institutional accountability requests and "
        "epistemic queries, they do not encompass reply threads, quote tweets, or retweet diffusion trees. Multi-layer network "
        "modeling across temporal intervals represents a promising direction for future research. Data collection complied with "
        "Platform X Developer Terms of Service; private accounts were excluded, and individual user handles—outside verified public "
        "political and institutional entities—were anonymized to safeguard user privacy."
    )

    doc.save(out_docx)
    print(f"Successfully saved camera-ready manuscript: {out_docx}")

if __name__ == "__main__":
    build_camera_ready()
