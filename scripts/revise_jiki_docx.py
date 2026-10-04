import os
import shutil
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def apply_table_styles(tbl, caption):
    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="808080"/>'
        '<w:left w:val="none"/>'
        '<w:right w:val="none"/>'
        '<w:insideV w:val="none"/>'
        '</w:tblBorders>'
    )
    tblPr.append(borders)
    if caption:
        cap_el = parse_xml(f'<w:tblCaption {nsdecls("w")} w:val="{caption}"/>')
        tblPr.append(cap_el)

def set_cell_properties(cell, fill_hex=None, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    if fill_hex:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_table(doc, headers, data, caption_title):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    apply_table_styles(tbl, caption_title)
    
    # Format Header Row
    hdr_cells = tbl.rows[0].cells
    for col_idx, header_text in enumerate(headers):
        hdr_cells[col_idx].text = header_text
        set_cell_properties(hdr_cells[col_idx], fill_hex="F2F2F2")
        for p in hdr_cells[col_idx].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.bold = True
                
    # Format Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_properties(row_cells[col_idx])
            for p in row_cells[col_idx].paragraphs:
                # Left align first column, center others
                if col_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(9.0)
                    if "Macro Average" in str(row_data[0]) or "Weighted Average" in str(row_data[0]) or "IndoBERT" in str(row_data[0]):
                        r.font.bold = True
                        
    return tbl

print("Table helper functions defined successfully")
