from pathlib import Path
import shutil
import zipfile
import re
import os
import sys

BASE = Path("/Users/jevin/Documents/tesis_mbg/backend")
OUT = BASE / "telegram_literature_output"
DESKTOP = Path("/Users/jevin/Desktop/TELEGRAM_LITERATURE_FINAL")

DESKTOP.mkdir(parents=True, exist_ok=True)

print("=" * 70)
print("FINAL TELEGRAM LITERATURE OUTPUT AUDIT")
print("=" * 70)

if not OUT.exists():
    raise SystemExit(f"ERROR: folder output tidak ditemukan: {OUT}")

# ---------------------------------------------------------
# Load pandas
# ---------------------------------------------------------
try:
    import pandas as pd
except Exception as e:
    raise SystemExit(f"ERROR pandas: {e}")

# ---------------------------------------------------------
# Locate master dataset
# ---------------------------------------------------------
master_candidates = [
    OUT / "telegram_literature_master.csv",
    OUT / "12_telegram_literature_master.csv",
    BASE / "telegram_papers_clean.csv",
    BASE / "telegram_papers_latest.csv",
]

master = None
for p in master_candidates:
    if p.exists():
        master = p
        break

if master is None:
    raise SystemExit("ERROR: tidak ditemukan CSV master/clean dataset.")

print(f"\nMASTER DATA: {master}")

df = pd.read_csv(master)
print(f"Jumlah baris: {len(df):,}")
print(f"Jumlah kolom: {len(df.columns):,}")

# ---------------------------------------------------------
# Utility column finder
# ---------------------------------------------------------
def find_col(keys):
    for c in df.columns:
        cl = str(c).lower()
        if any(k in cl for k in keys):
            return c
    return None

title_col = find_col(["title", "judul"])
doi_col = find_col(["doi"])
relevance_col = find_col(["relevance_score", "relevance"])
category_col = find_col(["relevance_category", "category", "kategori"])

print(f"Title column      : {title_col}")
print(f"DOI column        : {doi_col}")
print(f"Relevance column  : {relevance_col}")
print(f"Category column   : {category_col}")

# ---------------------------------------------------------
# DOI audit
# ---------------------------------------------------------
doi_missing = 0
doi_invalid = 0
doi_valid = 0

doi_re = re.compile(
    r"^10\.\d{4,9}/[-._;()/:A-Z0-9]+$",
    re.IGNORECASE
)

if doi_col:
    for value in df[doi_col].fillna("").astype(str):
        v = value.strip()
        if not v or v.lower() in {"nan", "none", "null", "n/a"}:
            doi_missing += 1
        elif doi_re.match(v):
            doi_valid += 1
        else:
            doi_invalid += 1

print("\nDOI AUDIT")
print(f"Valid syntax       : {doi_valid:,}")
print(f"Missing DOI        : {doi_missing:,}")
print(f"Invalid DOI syntax : {doi_invalid:,}")

# ---------------------------------------------------------
# Scopus / Q1 audit
# IMPORTANT:
# We do NOT convert "candidate" into "verified".
# ---------------------------------------------------------
scopus_cols = [
    c for c in df.columns
    if "scopus" in str(c).lower()
]

q1_cols = [
    c for c in df.columns
    if "q1" in str(c).lower() or "quartile" in str(c).lower()
]

def collect_values(cols):
    values = {}
    for c in cols:
        vals = (
            df[c]
            .fillna("UNVERIFIED")
            .astype(str)
            .str.strip()
            .replace("", "UNVERIFIED")
            .value_counts()
            .to_dict()
        )
        values[c] = vals
    return values

scopus_values = collect_values(scopus_cols)
q1_values = collect_values(q1_cols)

print("\nSCOPUS COLUMNS")
if scopus_values:
    for c, vals in scopus_values.items():
        print(f"{c}: {vals}")
else:
    print("Tidak ada kolom Scopus.")

print("\nQ1 / QUARTILE COLUMNS")
if q1_values:
    for c, vals in q1_values.items():
        print(f"{c}: {vals}")
else:
    print("Tidak ada kolom Q1/Quartile.")

# ---------------------------------------------------------
# Detect possible fabricated/suspicious status labels
# ---------------------------------------------------------
warning_terms = {
    "VERIFIED",
    "VERIFIED_SCOPUS",
    "SCOPUS_VERIFIED",
    "Q1_VERIFIED",
    "VERIFIED_Q1",
}

suspicious_statuses = []

for c in scopus_cols + q1_cols:
    vals = df[c].fillna("").astype(str).str.upper().str.strip()
    for v in vals.unique():
        if v in warning_terms:
            suspicious_statuses.append((c, v))

# We do not alter the source data.
# We simply flag it in the audit.
if suspicious_statuses:
    print("\nWARNING: ditemukan label verifikasi yang perlu diperiksa sumbernya:")
    for x in suspicious_statuses:
        print(" ", x)

# ---------------------------------------------------------
# Relevance statistics
# ---------------------------------------------------------
relevance_summary = {}
if category_col:
    relevance_summary = (
        df[category_col]
        .fillna("UNCLASSIFIED")
        .astype(str)
        .value_counts()
        .to_dict()
    )

print("\nRELEVANCE")
if relevance_summary:
    for k, v in relevance_summary.items():
        print(f"{k}: {v:,}")
else:
    print("Kolom kategori relevance tidak ditemukan.")

# ---------------------------------------------------------
# Read pipeline error log
# ---------------------------------------------------------
error_log = OUT / "pipeline_errors.log"
error_text = ""

if error_log.exists():
    error_text = error_log.read_text(encoding="utf-8", errors="replace")
    print("\nPIPELINE ERROR LOG")
    if error_text.strip():
        print(error_text[-5000:])
    else:
        print("Kosong.")
else:
    print("\nPIPELINE ERROR LOG: tidak ada.")

# ---------------------------------------------------------
# Build final audit report
# ---------------------------------------------------------
report = OUT / "FINAL_PIPELINE_REPORT.txt"

files_before = sorted(
    p for p in OUT.iterdir()
    if p.is_file()
)

with report.open("w", encoding="utf-8") as f:
    f.write("FINAL TELEGRAM LITERATURE PIPELINE REPORT\n")
    f.write("=" * 70 + "\n\n")

    f.write(f"Backend : {BASE}\n")
    f.write(f"Output  : {OUT}\n")
    f.write(f"Master  : {master}\n\n")

    f.write("DATASET\n")
    f.write(f"- Records : {len(df):,}\n")
    f.write(f"- Columns : {len(df.columns):,}\n\n")

    f.write("DOI AUDIT\n")
    f.write(f"- Valid syntax       : {doi_valid:,}\n")
    f.write(f"- Missing DOI        : {doi_missing:,}\n")
    f.write(f"- Invalid DOI syntax : {doi_invalid:,}\n\n")

    f.write("SCOPUS / Q1 INTEGRITY RULE\n")
    f.write("- No Scopus or Q1 status is invented by this final packaging step.\n")
    f.write("- Candidate/possible labels are NOT treated as verified.\n")
    f.write("- Information that cannot be independently verified remains UNVERIFIED.\n\n")

    f.write("SCOPUS COLUMNS\n")
    if scopus_values:
        for c, vals in scopus_values.items():
            f.write(f"{c}: {vals}\n")
    else:
        f.write("No Scopus column detected.\n")

    f.write("\nQ1 / QUARTILE COLUMNS\n")
    if q1_values:
        for c, vals in q1_values.items():
            f.write(f"{c}: {vals}\n")
    else:
        f.write("No Q1/Quartile column detected.\n")

    f.write("\nRELEVANCE\n")
    for k, v in relevance_summary.items():
        f.write(f"- {k}: {v:,}\n")

    f.write("\nPIPELINE ERROR LOG\n")
    if error_text.strip():
        f.write(error_text[-10000:])
    else:
        f.write("No recorded pipeline error.\n")

    f.write("\nOUTPUT FILES\n")
    for p in files_before:
        f.write(f"- {p.name} | {p.stat().st_size:,} bytes\n")

# ---------------------------------------------------------
# Create Word document
# ---------------------------------------------------------
try:
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
except Exception as e:
    raise SystemExit(
        "python-docx belum tersedia. "
        f"Error: {e}"
    )

doc = Document()

section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("FINAL TELEGRAM LITERATURE RESEARCH")
r.bold = True
r.font.name = "Times New Roman"
r.font.size = Pt(16)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(
    "Literature pipeline 2024–2026 | Audit & evidence integrity report"
)
r.italic = True
r.font.name = "Times New Roman"
r.font.size = Pt(10)

doc.add_paragraph("")

# Summary
doc.add_heading("1. Dataset Summary", level=1)

summary_rows = [
    ("Master dataset", str(master)),
    ("Total records", f"{len(df):,}"),
    ("Total columns", f"{len(df.columns):,}"),
    ("Valid DOI syntax", f"{doi_valid:,}"),
    ("Missing DOI", f"{doi_missing:,}"),
    ("Invalid DOI syntax", f"{doi_invalid:,}"),
]

table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"

hdr = table.rows[0].cells
hdr[0].text = "Metric"
hdr[1].text = "Result"

for a, b in summary_rows:
    cells = table.add_row().cells
    cells[0].text = a
    cells[1].text = b

# Scopus / Q1
doc.add_heading("2. Scopus and Q1 Verification Status", level=1)

p = doc.add_paragraph()
p.add_run(
    "Integrity rule: "
).bold = True
p.add_run(
    "Scopus/Q1 is not considered verified merely because a paper is a candidate. "
    "Where independent verification is unavailable, the status is treated as UNVERIFIED."
)

if scopus_values:
    for c, vals in scopus_values.items():
        doc.add_paragraph(f"{c}: {vals}")
else:
    doc.add_paragraph("No Scopus verification column detected.")

if q1_values:
    for c, vals in q1_values.items():
        doc.add_paragraph(f"{c}: {vals}")
else:
    doc.add_paragraph("No Q1/quartile verification column detected.")

if suspicious_statuses:
    doc.add_paragraph(
        "AUDIT FLAG: the source dataset contains verification-style labels. "
        "Those labels were not independently upgraded by this script."
    )

# Relevance
doc.add_heading("3. Relevance Screening", level=1)

if relevance_summary:
    t = doc.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].text = "Category"
    t.rows[0].cells[1].text = "Records"

    for k, v in relevance_summary.items():
        cells = t.add_row().cells
        cells[0].text = str(k)
        cells[1].text = f"{v:,}"
else:
    doc.add_paragraph("No relevance category column available.")

# Candidate papers
doc.add_heading("4. Literature Records", level=1)

candidate = df.copy()

if relevance_col:
    candidate["_sort_relevance"] = pd.to_numeric(
        candidate[relevance_col], errors="coerce"
    )
    candidate = candidate.sort_values(
        "_sort_relevance", ascending=False, na_position="last"
    )

candidate = candidate.head(40)

cols = []
for c in [title_col, doi_col, category_col, relevance_col]:
    if c and c not in cols:
        cols.append(c)

if not cols:
    cols = list(df.columns[:4])

t = doc.add_table(rows=1, cols=len(cols))
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, c in enumerate(cols):
    t.rows[0].cells[i].text = str(c)

for _, row in candidate.iterrows():
    cells = t.add_row().cells
    for i, c in enumerate(cols):
        value = row.get(c, "")
        if pd.isna(value):
            value = ""
        cells[i].text = str(value)[:500]

# Error log
doc.add_heading("5. Pipeline Error Log", level=1)

if error_text.strip():
    p = doc.add_paragraph()
    p.add_run(
        "The pipeline recorded the following messages. "
        "This section is diagnostic and does not alter the source dataset."
    )
    doc.add_paragraph(error_text[-10000:])
else:
    doc.add_paragraph("No pipeline error was recorded.")

# Output list
doc.add_heading("6. Generated Output Files", level=1)

for p in sorted(OUT.iterdir()):
    if p.is_file():
        doc.add_paragraph(
            f"{p.name} — {p.stat().st_size:,} bytes"
        )

docx_path = OUT / "FINAL_TELEGRAM_LITERATURE_RESEARCH.docx"
doc.save(docx_path)

print(f"\nWORD CREATED: {docx_path}")
print(f"SIZE: {docx_path.stat().st_size:,} bytes")

# ---------------------------------------------------------
# Copy every output to Desktop package
# ---------------------------------------------------------
for old in DESKTOP.iterdir():
    if old.is_file():
        old.unlink()
    elif old.is_dir():
        shutil.rmtree(old)

for p in OUT.iterdir():
    if p.is_file():
        shutil.copy2(p, DESKTOP / p.name)

# Add report explicitly if somehow missing
if report.exists():
    shutil.copy2(report, DESKTOP / report.name)

# ---------------------------------------------------------
# Create ZIP on Desktop
# ---------------------------------------------------------
zip_path = Path("/Users/jevin/Desktop/FINAL_TELEGRAM_LITERATURE_FINAL.zip")

if zip_path.exists():
    zip_path.unlink()

with zipfile.ZipFile(
    zip_path,
    "w",
    compression=zipfile.ZIP_DEFLATED
) as z:
    for p in sorted(DESKTOP.iterdir()):
        if p.is_file():
            z.write(p, arcname=p.name)

print(f"ZIP CREATED : {zip_path}")
print(f"ZIP SIZE    : {zip_path.stat().st_size:,} bytes")

# ---------------------------------------------------------
# FINAL STATISTICS
# ---------------------------------------------------------
desktop_files = sorted(
    p for p in DESKTOP.iterdir()
    if p.is_file()
)

print("\n" + "=" * 70)
print("FINAL STATISTICS")
print("=" * 70)
print(f"Records                    : {len(df):,}")
print(f"Valid DOI syntax            : {doi_valid:,}")
print(f"Missing DOI                 : {doi_missing:,}")
print(f"Invalid DOI syntax          : {doi_invalid:,}")
print(f"Scopus columns              : {len(scopus_cols)}")
print(f"Q1/Quartile columns         : {len(q1_cols)}")
print(f"Suspicious verification lbl : {len(suspicious_statuses)}")
print(f"Desktop output files        : {len(desktop_files)}")

print("\n" + "=" * 70)
print("ABSOLUTE OUTPUT PATHS")
print("=" * 70)

for p in desktop_files:
    print(f"{p} | {p.stat().st_size:,} bytes")

print(f"\nWORD : {docx_path}")
print(f"ZIP  : {zip_path}")

print("\n" + "=" * 70)
print("OPENING FOLDER")
print("=" * 70)

os.system(
    'open "/Users/jevin/Desktop/TELEGRAM_LITERATURE_FINAL"'
)
